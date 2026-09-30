"""Génération de la note de comité (texte Markdown et fichier Word)."""
from __future__ import annotations

import io

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from src import data as D

COUL = {"Rouge": "D64545", "Orange": "E8A317", "Vert": "2E9E5B"}


def contenu(ctx: dict) -> dict:
    ev, risques, prev = ctx["ev"], ctx["risques"], ctx["prev"]
    top = ev[ev["statut"] != "Vert"].sort_values("score_priorite", ascending=False).head(3)
    decisions = []
    for r in top.itertuples():
        action = r.action_prioritaire if isinstance(r.action_prioritaire, str) else "définir un plan d'action"
        p = prev.set_index("code").loc[r.code]
        fin = "" if p["lance"] else f" Lancement estimé au {D.fmt_date(p['pire_cas'])} (prévu le {D.fmt_date(p['fin_prevue'])})."
        decisions.append(
            f"{r.code} – {r.produit} ({r.statut}) : SPI {D.fmt_idx(r.spi)}, CPI {D.fmt_idx(r.cpi)}, "
            f"dépassement attendu {D.fmt_m(max(0, -r.vac))}.{fin} Décision demandée : {action.lower()}.")
    retard = risques[risques["action_en_retard"]]
    return {
        "titre": f"Note au comité de pilotage — situation au {D.fmt_date(ctx['date'])}",
        "messages": D.messages_cles(ev, risques),
        "decisions": decisions,
        "actions_retard": [f"{x.id} ({x.produit}) : {x.action} — {x.responsable}, échue depuis "
                           f"{int(x.jours_depassement)} jours" for x in retard.itertuples()],
        "tableau": ev[["code", "produit", "statut", "spi", "cpi", "budget", "eac", "vac"]],
    }


def markdown(c: dict) -> str:
    lignes = [f"# {c['titre']}", "", "## Synthèse"] + [f"- {m}" for m in c["messages"]]
    lignes += ["", "## Décisions demandées au comité"] + [f"{i}. {d}" for i, d in enumerate(c["decisions"], 1)]
    if c["actions_retard"]:
        lignes += ["", "## Actions correctives à relancer"] + [f"- {a}" for a in c["actions_retard"]]
    lignes += ["", "## Situation par produit", "",
               "| Code | Produit | Statut | SPI | CPI | Budget (M) | Coût final estimé (M) | Écart (M) |",
               "|---|---|---|---|---|---|---|---|"]
    for r in c["tableau"].itertuples():
        lignes.append(f"| {r.code} | {r.produit} | {r.statut} | {D.fmt_idx(r.spi)} | {D.fmt_idx(r.cpi)} | "
                      f"{D.fmt_nb(r.budget)} | {D.fmt_nb(r.eac)} | {D.fmt_nb(r.vac)} |")
    lignes += ["", "_Données synthétiques — projet de démonstration._"]
    return "\n".join(lignes)


def _fond(cell, hexa: str) -> None:
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexa)
    tc.append(shd)


def word(c: dict) -> bytes:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    h = doc.add_heading(c["titre"], level=1)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    doc.add_heading("Synthèse", level=2)
    for m in c["messages"]:
        doc.add_paragraph(m, style="List Bullet")
    doc.add_heading("Décisions demandées au comité", level=2)
    for d in c["decisions"]:
        doc.add_paragraph(d, style="List Number")
    if c["actions_retard"]:
        doc.add_heading("Actions correctives à relancer", level=2)
        for a in c["actions_retard"]:
            doc.add_paragraph(a, style="List Bullet")
    doc.add_heading("Situation par produit", level=2)
    entetes = ["Code", "Produit", "Statut", "SPI", "CPI", "Budget (M)", "Coût final (M)", "Écart (M)"]
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Light Grid Accent 1"
    for i, e in enumerate(entetes):
        t.rows[0].cells[i].text = e
    for r in c["tableau"].itertuples():
        cells = t.add_row().cells
        vals = [r.code, r.produit, r.statut, D.fmt_idx(r.spi), D.fmt_idx(r.cpi),
                D.fmt_nb(r.budget), D.fmt_nb(r.eac), D.fmt_nb(r.vac)]
        for i, v in enumerate(vals):
            cells[i].text = str(v)
            if i >= 3:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _fond(cells[2], COUL.get(r.statut, "FFFFFF"))
        cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p = doc.add_paragraph("Données synthétiques — projet de démonstration.")
    p.runs[0].italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def excel(ctx: dict) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        cols = {"code": "Code", "produit": "Produit", "budget": "Budget", "pv": "PV", "ev": "EV", "ac": "AC",
                "spi": "SPI", "cpi": "CPI", "eac": "EAC", "vac": "VAC", "tcpi": "TCPI", "statut": "Statut",
                "motif": "Motif", "score_priorite": "Score priorité"}
        ctx["ev"][list(cols)].rename(columns=cols).to_excel(w, sheet_name="Valeur acquise", index=False)
        ctx["prev"].drop(columns=["lance"]).to_excel(w, sheet_name="Prévisions", index=False)
        ctx["risques"].to_excel(w, sheet_name="Risques", index=False)
        ctx["data"]["jalons"].to_excel(w, sheet_name="Jalons", index=False)
    return buf.getvalue()
